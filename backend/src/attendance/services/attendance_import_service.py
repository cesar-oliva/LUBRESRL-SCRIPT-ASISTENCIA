from datetime import date, datetime, timedelta

from src.attendance.models.turn import Turn, WorkPeriod
from src.attendance.repositories.attendance_report_repository import AttendanceReportRepository
from src.attendance.repositories.employee_repository import EmployeeRepository
from src.attendance.repositories.employee_turn_repository import EmployeeTurnRepository
from src.attendance.repositories.turn_repository import TurnRepository
from src.attendance.repositories.monthly_turn_assignment_repository import (
    MonthlyTurnAssignmentRepository,
)
from src.attendance.repositories.medical_certificate_repository import (
    MedicalCertificateRepository,
)


class AttendanceImportService:
    VALID_REGISTRATIONS = {0, 1}

    def __init__(
        self,
        employee_repository=None,
        employee_turn_repository=None,
        turn_repository=None,
        report_repository=None,
        monthly_turn_assignment_repository=None,
        medical_certificate_repository=None,
    ):
        self.employee_repository = employee_repository or EmployeeRepository()
        self.employee_turn_repository = employee_turn_repository or EmployeeTurnRepository()
        self.turn_repository = turn_repository or TurnRepository()
        self.report_repository = report_repository or AttendanceReportRepository()
        self.monthly_turn_assignment_repository = (
            monthly_turn_assignment_repository or MonthlyTurnAssignmentRepository()
        )
        self.medical_certificate_repository = (
            medical_certificate_repository or MedicalCertificateRepository()
        )

    def read_excel_rows(self, file_obj):
        from io import BytesIO
        from openpyxl import load_workbook

        if isinstance(file_obj, (bytes, bytearray)):
            workbook = load_workbook(BytesIO(file_obj), read_only=True, data_only=True)
        else:
            workbook = load_workbook(file_obj, read_only=True, data_only=True)

        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return [], []

        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        required = ["Usuario Nro.", "Fecha/Hora", "Registro"]
        missing = [item for item in required if item not in headers]
        if missing:
            return [], [f"Faltan columnas requeridas: {', '.join(missing)}"]

        indexes = {name: headers.index(name) for name in required}
        normalized = []
        errors = []

        for row_number, row in enumerate(rows[1:], start=2):
            if not row or not any(cell is not None and str(cell).strip() != "" for cell in row):
                continue

            try:
                employee_number = int(str(row[indexes["Usuario Nro."]]).strip())
            except (TypeError, ValueError):
                errors.append(f"Fila {row_number}: legajo inválido.")
                continue

            try:
                dt = self._parse_datetime(row[indexes["Fecha/Hora"]])
            except ValueError as exc:
                errors.append(f"Fila {row_number}: {exc}")
                continue

            registration = self._normalize_registration(row[indexes["Registro"]])
            if registration is None:
                errors.append(
                    f"Fila {row_number}: registro inválido ({row[indexes['Registro']]}); "
                    "solo se admiten entradas/salidas compatibles."
                )
                continue

            if not self.employee_repository.exists(employee_number):
                errors.append(f"Fila {row_number}: el legajo {employee_number} no existe en la base.")
                continue

            normalized.append({
                "employee_number": employee_number,
                "date_time": dt,
                "registration": registration,
            })

        sorted_rows, duplicated = self._sort_and_deduplicate(normalized)
        return sorted_rows, errors + (["Se detectaron registros duplicados exactos."] if duplicated else [])

    def process_rows(self, rows, period, tolerance_minutes=5):
        if rows is None:
            rows = []

        normalized = []
        errors = []
        for row in rows:
            try:
                employee_number = int(row["employee_number"])
                dt = self._parse_datetime(row["date_time"])
                registration = self._normalize_registration(row["registration"])
            except Exception as exc:
                errors.append(f"Registro inválido: {exc}")
                continue

            if registration not in self.VALID_REGISTRATIONS:
                errors.append(f"Legajo {employee_number}: registro inválido ({registration}).")
                continue

            normalized.append({
                "employee_number": employee_number,
                "date_time": dt,
                "registration": registration,
            })

        normalized, duplicated_records = self._sort_and_deduplicate(normalized)
        grouped = self._group_by_employee_day(normalized)
        summary = {
            "total_employees": len({item["employee_number"] for item in normalized}),
            "total_turns_analyzed": 0,
            "total_imported": 0,
            "duplicated_records": duplicated_records,
            "late_arrivals": 0,
            "reincidences": 0,
            "early_departures": 0,
            "missing_records": 0,
            "absence_records": 0,
            "medical_absences": 0,
            "unjustified_absences": 0,
            "inconsistent_records": 0,
        }

        entries = []
        period_days = self._period_days(period)
        certificates = self._active_certificates_by_employee(period)
        for employee_number, day_key in self._missing_assigned_days(grouped, period_days):
            grouped[(employee_number, day_key)] = []

        for (employee_number, day_key), records in sorted(grouped.items(), key=lambda item: (item[0][0], item[0][1])):
            employee = self.employee_repository.get_by_employee_number(employee_number)
            segments = self._build_employee_segments(records) if records else [{"entry": None, "exit": None}]
            day_turn = self._resolve_turn_for_segments(
                employee_number,
                day_key,
                segments,
                period=period,
                tolerance_minutes=tolerance_minutes,
            ) if employee else None

            for segment_index, segment in enumerate(segments):
                actual_entry = segment["entry"]["date_time"] if segment["entry"] else None
                actual_exit = segment["exit"]["date_time"] if segment["exit"] else None

                if employee is None:
                    turn = None
                    expected_entry = None
                    expected_exit = None
                    minutes_late = 0
                    minutes_early = 0
                    status = "REGISTRO_INCONSISTENTE"
                    observation = "El legajo no existe en la base de datos."
                else:
                    turn = day_turn or self._resolve_turn_for_day(
                        employee_number,
                        day_key,
                        actual_entry,
                        actual_exit,
                        period=period,
                        tolerance_minutes=tolerance_minutes,
                    )
                    expected_entry, expected_exit = self._resolve_expected_interval_for_segment(
                        turn,
                        day_key,
                        actual_entry,
                        actual_exit,
                        segment_index,
                    )
                    minutes_late = self._compute_minutes_late(actual_entry, expected_entry, tolerance_minutes)
                    minutes_early = self._compute_minutes_early(actual_exit, expected_exit, tolerance_minutes)

                    if actual_entry and actual_exit:
                        if minutes_late > 0 and minutes_early > 0:
                            status = "LLEGADA_TARDE_Y_SALIDA_ANTICIPADA"
                            observation = "Entrada registrada después del horario establecido y salida registrada antes del horario establecido."
                        elif minutes_late > 0:
                            status = "LLEGADA_TARDE"
                            observation = f"Entrada registrada {minutes_late} minutos después del horario establecido."
                        elif minutes_early > 0:
                            status = "SALIDA_ANTICIPADA"
                            observation = f"Salida registrada {minutes_early} minutos antes del horario establecido."
                        else:
                            status = "EN_HORARIO"
                            observation = "Marcación dentro del horario establecido."
                    elif actual_entry and not actual_exit:
                        status = "SIN_REGISTRO_SALIDA"
                        observation = "No se encontró registro de salida."
                    elif not actual_entry and actual_exit:
                        status = "SIN_REGISTRO_ENTRADA"
                        observation = "No se encontró registro de entrada."
                    else:
                        if day_key in certificates.get(employee_number, set()):
                            status = "FALTA_CERTIFICADO"
                            observation = "Certificado validado en centro médico."
                        else:
                            status = "FALTA_SIN_CERTIFICADO"
                            observation = "Falta sin certificado médico activo."

                    if period and not self._is_in_period(day_key, period):
                        status = "REGISTRO_INCONSISTENTE"
                        observation = "El registro se encuentra fuera del período seleccionado."

                output = {
                    "employee_number": employee_number,
                    "employee_name": employee.name if employee else "",
                    "date": day_key.strftime("%Y-%m-%d"),
                    "turn": turn.code if turn and hasattr(turn, "code") else None,
                    "expected_entry": expected_entry.strftime("%H:%M") if expected_entry else None,
                    "actual_entry": actual_entry.strftime("%Y-%m-%d %H:%M") if actual_entry else None,
                    "expected_exit": expected_exit.strftime("%H:%M") if expected_exit else None,
                    "actual_exit": actual_exit.strftime("%Y-%m-%d %H:%M") if actual_exit else None,
                    "status": status,
                    "observation": observation,
                    "minutes_late": minutes_late if employee else 0,
                    "minutes_early": minutes_early if employee else 0,
                    "entry_after_expected": bool(
                        actual_entry and expected_entry and actual_entry > expected_entry
                    ),
                    "state": status,
                }
                entries.append(output)
                summary["total_imported"] += 1

                if minutes_late > 0:
                    summary["late_arrivals"] += 1
                if minutes_early > 0:
                    summary["early_departures"] += 1
                if status in {"SIN_REGISTRO", "SIN_REGISTRO_ENTRADA", "SIN_REGISTRO_SALIDA"}:
                    summary["missing_records"] += 1
                if status in {"FALTA_SIN_CERTIFICADO", "FALTA_CERTIFICADO"}:
                    summary["missing_records"] += 1
                    summary["absence_records"] += 1
                if status == "FALTA_CERTIFICADO":
                    summary["medical_absences"] += 1
                if status == "FALTA_SIN_CERTIFICADO":
                    summary["unjustified_absences"] += 1
                if status == "REGISTRO_INCONSISTENTE":
                    summary["inconsistent_records"] += 1

        late_by_employee = {}
        for entry in entries:
            if entry["entry_after_expected"]:
                late_by_employee[entry["employee_number"]] = (
                    late_by_employee.get(entry["employee_number"], 0) + 1
                )

        recurring_employees = {
            employee_number
            for employee_number, count in late_by_employee.items()
            if count >= 3
        }
        summary["reincidences"] = len(recurring_employees)
        for entry in entries:
            if (
                entry["employee_number"] not in recurring_employees
                or not entry["entry_after_expected"]
            ):
                continue
            entry["status"] = f"REINCIDENCIA_{entry['status']}"
            entry["state"] = entry["status"]

        summary["total_turns_analyzed"] = len(entries)
        summary["total_employees"] = len({entry["employee_number"] for entry in entries})

        if entries and summary["inconsistent_records"] == len(entries):
            errors.append(
                "Todos los registros quedaron fuera del período seleccionado. "
                "Revise el período (YYYY-MM) antes de confirmar la importación."
            )

        return {"entries": entries, "summary": summary, "errors": errors}

    def _resolve_turn_for_segments(
        self,
        employee_number,
        day_key,
        segments,
        period=None,
        tolerance_minutes=5,
    ):
        monthly_code = self.monthly_turn_assignment_repository.get_code(
            period,
            day_key,
            employee_number,
        ) if period else None
        assigned_turn = self._get_turn_by_code(monthly_code) if monthly_code else None

        candidates = self.turn_repository.get_active() or self.turn_repository.get_all()
        matching_turn = self._find_turn_matching_segments(
            candidates,
            segments,
            day_key,
            tolerance_minutes,
        )
        if matching_turn is not None:
            if (
                period
                and monthly_code
                and self._normalize_turn_code(matching_turn.code) != self._normalize_turn_code(monthly_code)
            ):
                self.monthly_turn_assignment_repository.save_code(
                    period,
                    day_key,
                    employee_number,
                    matching_turn.code,
                )
            return matching_turn

        if assigned_turn is not None and assigned_turn.active:
            return assigned_turn

        weekday = day_key.isoweekday()
        schedule = self.employee_turn_repository.get_employee_schedule(employee_number)
        for item in schedule:
            if item["day_of_week"] == weekday:
                return item["turn"]
        return self._resolve_turn_for_day(
            employee_number,
            day_key,
            segments[0]["entry"]["date_time"] if segments and segments[0]["entry"] else None,
            segments[0]["exit"]["date_time"] if segments and segments[0]["exit"] else None,
            period=period,
            tolerance_minutes=tolerance_minutes,
        )

    def _find_turn_matching_segments(self, turns, segments, day_key, tolerance_minutes):
        matches = []
        for turn in turns or []:
            if not turn.active or len(getattr(turn, "periods", [])) != len(segments):
                continue
            score = 0
            matched = True
            for segment, work_period in zip(segments, turn.periods):
                expected_entry = datetime.combine(day_key, self._parse_time(work_period.start_time))
                expected_exit = datetime.combine(day_key, self._parse_time(work_period.end_time))
                if expected_exit < expected_entry:
                    expected_exit += timedelta(days=1)
                actual_entry = segment["entry"]["date_time"] if segment["entry"] else None
                actual_exit = segment["exit"]["date_time"] if segment["exit"] else None
                if not actual_entry or not actual_exit:
                    matched = False
                    break
                adjusted_exit = actual_exit
                if adjusted_exit < actual_entry:
                    adjusted_exit += timedelta(days=1)
                entry_delta = abs((actual_entry - expected_entry).total_seconds())
                exit_delta = abs((adjusted_exit - expected_exit).total_seconds())
                if entry_delta > 30 * 60 or exit_delta > 30 * 60:
                    matched = False
                    break
                score += entry_delta + exit_delta
            if matched:
                matches.append((score, turn))
        return min(matches, key=lambda item: item[0])[1] if matches else None

    @staticmethod
    def _normalize_turn_code(code):
        return str(code or "").replace(" ", "").upper()

    def _missing_assigned_days(self, grouped, period_days):
        missing = []
        monthly_assignments = self.monthly_turn_assignment_repository.get_for_period(
            self._period_key(period_days)
        ) if period_days else []
        monthly_days = {
            (row["employee_number"], self._as_date(row["assignment_date"]))
            for row in monthly_assignments
            if row["assignment_code"] == "42"
            or self._get_turn_by_code(row["assignment_code"]) is not None
        }
        for employee in self.employee_repository.get_all():
            if not employee.active:
                continue
            schedule = self.employee_turn_repository.get_employee_schedule(employee.employee_number)
            assigned_weekdays = {
                item["day_of_week"]
                for item in schedule
                if item.get("turn") and item["turn"].active
            }
            for day_key in period_days:
                key = (employee.employee_number, day_key)
                if (
                    (day_key.isoweekday() in assigned_weekdays or key in monthly_days)
                    and key not in grouped
                ):
                    missing.append((employee.employee_number, day_key))
        return missing

    @staticmethod
    def _period_key(period_days):
        return period_days[0].strftime("%Y-%m")

    def _get_turn_by_code(self, code):
        turn = self.turn_repository.get_by_code(code)
        if turn is not None:
            return turn

        normalized_code = str(code).replace(" ", "").upper()
        for candidate in self.turn_repository.get_all() or []:
            if str(candidate.code).replace(" ", "").upper() == normalized_code:
                return candidate
        return None

    def _active_certificates_by_employee(self, period):
        certificates_by_employee = {}
        for certificate in self.medical_certificate_repository.get_active_covering_period(period):
            day = self._as_date(certificate.valid_from)
            valid_until = self._as_date(certificate.valid_until)
            while day <= valid_until:
                if self._is_in_period(day, period):
                    certificates_by_employee.setdefault(certificate.employee_number, set()).add(day)
                day += timedelta(days=1)
        return certificates_by_employee

    @staticmethod
    def _as_date(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return datetime.strptime(value, "%Y-%m-%d").date()

    def _period_days(self, period):
        try:
            first_day = datetime.strptime(period + "-01", "%Y-%m-%d").date()
        except (TypeError, ValueError):
            return []
        next_month = (first_day.replace(day=28) + timedelta(days=4)).replace(day=1)
        return [
            first_day + timedelta(days=offset)
            for offset in range((next_month - first_day).days)
        ]

    def save_report(self, period, entry):
        employee = self.employee_repository.get_by_employee_number(entry["employee_number"])
        report_data = {
            "period": period,
            "employee_number": entry["employee_number"],
            "employee_name": employee.name if employee else entry.get("employee_name", ""),
            "report_date": entry["date"],
            "turn_code": entry.get("turn"),
            "expected_entry": entry.get("expected_entry"),
            "actual_entry": entry.get("actual_entry"),
            "expected_exit": entry.get("expected_exit"),
            "actual_exit": entry.get("actual_exit"),
            "status": entry.get("status") or entry.get("state"),
            "observation": entry.get("observation"),
            "minutes_late": entry.get("minutes_late", 0),
            "minutes_early": entry.get("minutes_early", 0),
            "active": True,
        }
        report = type("Report", (), report_data)()
        return self.report_repository.save(report)

    def get_report(self, period):
        return self.report_repository.get_by_period(period)

    def replace_period_report(self, period, entries):
        self.report_repository.delete_by_period(period)
        saved = []
        for entry in entries:
            saved.append(self.save_report(period, entry))
        return saved

    def update_report_observation(self, report_id, observation):
        updated = self.report_repository.update_observation(report_id, observation)
        if not updated:
            return None
        return self.report_repository.get_by_id(report_id)

    def _parse_datetime(self, value):
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        if isinstance(value, str):
            cleaned = value.strip()
            for fmt in [
                "%d/%m/%Y %H:%M",
                "%d/%m/%Y %H:%M:%S",
                "%Y-%m-%d %H:%M",
                "%Y-%m-%d %H:%M:%S",
                "%d-%m-%Y %H:%M",
                "%d-%m-%Y %H:%M:%S",
            ]:
                try:
                    return datetime.strptime(cleaned, fmt)
                except ValueError:
                    continue
            raise ValueError(f"No se pudo convertir '{value}' a datetime válido.")
        raise ValueError(f"Tipo de dato no soportado para fecha/hora: {type(value)}")

    def _sort_and_deduplicate(self, rows):
        ordered = sorted(rows, key=lambda item: (item["employee_number"], item["date_time"], item["registration"]))
        seen = set()
        unique = []
        duplicated = 0

        for row in ordered:
            unique_key = (row["employee_number"], row["date_time"])
            if unique_key in seen:
                duplicated += 1
                continue
            seen.add(unique_key)
            unique.append(row)

        return unique, duplicated

    def _group_by_employee_day(self, rows):
        grouped = {}
        for row in rows:
            key = (row["employee_number"], row["date_time"].date())
            grouped.setdefault(key, []).append(row)
        return grouped

    def _resolve_turn_for_day(
        self,
        employee_number,
        day_key,
        actual_entry=None,
        actual_exit=None,
        period=None,
        tolerance_minutes=5,
    ):
        monthly_code = None
        if period:
            monthly_code = self.monthly_turn_assignment_repository.get_code(
                period,
                day_key,
                employee_number,
            )

        if monthly_code:
            monthly_turn = self._get_turn_by_code(monthly_code)
            if monthly_turn is not None and monthly_turn.active:
                if self._entry_matches_turn(actual_entry, monthly_turn, day_key, tolerance_minutes):
                    return monthly_turn

                alternative = self._find_turn_matching_entry(
                    actual_entry,
                    actual_exit,
                    day_key,
                    tolerance_minutes,
                )
                if alternative is not None and alternative.code != monthly_turn.code:
                    self.monthly_turn_assignment_repository.save_code(
                        period,
                        day_key,
                        employee_number,
                        alternative.code,
                    )
                    return alternative

                return monthly_turn

        weekday = day_key.isoweekday()
        schedule = self.employee_turn_repository.get_employee_schedule(employee_number)
        for item in schedule:
            if item["day_of_week"] == weekday:
                return item["turn"]

        inferred_turn = self._infer_turn_from_marks(actual_entry, actual_exit)
        if inferred_turn is not None:
            return inferred_turn

        def anchor(dt):
            return dt.replace(minute=0, second=0, microsecond=0)

        if actual_entry and actual_exit:
            start_dt = anchor(actual_entry)
            end_dt = anchor(actual_exit)
            if end_dt < start_dt:
                end_dt += timedelta(days=1)
            return Turn(
                code="INF",
                name="Turno inferido",
                periods=[WorkPeriod(start_time=start_dt.strftime("%H:%M:%S"), end_time=end_dt.strftime("%H:%M:%S"))],
                active=True,
            )

        if actual_entry:
            start_dt = anchor(actual_entry)
            end_dt = start_dt + timedelta(hours=8)
            return Turn(
                code="INF",
                name="Turno inferido",
                periods=[WorkPeriod(start_time=start_dt.strftime("%H:%M:%S"), end_time=end_dt.strftime("%H:%M:%S"))],
                active=True,
            )

        if actual_exit:
            end_dt = anchor(actual_exit)
            start_dt = end_dt - timedelta(hours=8)
            return Turn(
                code="INF",
                name="Turno inferido",
                periods=[WorkPeriod(start_time=start_dt.strftime("%H:%M:%S"), end_time=end_dt.strftime("%H:%M:%S"))],
                active=True,
            )

        return None

    def _find_turn_matching_entry(self, actual_entry, actual_exit, day_key, tolerance_minutes):
        if not actual_entry:
            return None

        candidate = self._infer_turn_from_marks(actual_entry, actual_exit)
        if candidate is None or not self._entry_matches_turn(
            actual_entry,
            candidate,
            day_key,
            tolerance_minutes,
        ):
            return None
        return candidate

    def _entry_matches_turn(self, actual_entry, turn, day_key, tolerance_minutes):
        if not actual_entry or not turn or not getattr(turn, "periods", None):
            return False

        expected_entry = datetime.combine(
            day_key,
            self._parse_time(turn.periods[0].start_time),
        )
        difference = abs((actual_entry - expected_entry).total_seconds())
        if difference > 12 * 3600:
            expected_entry += timedelta(days=1 if actual_entry > expected_entry else -1)
            difference = abs((actual_entry - expected_entry).total_seconds())

        return difference <= tolerance_minutes * 60

    def _infer_turn_from_marks(self, actual_entry=None, actual_exit=None):
        turns = self.turn_repository.get_active() or self.turn_repository.get_all()
        if not turns:
            return None

        best_turn = None
        best_score = None

        for turn in turns:
            if not getattr(turn, "periods", None):
                continue

            first_period = turn.periods[0]
            last_period = turn.periods[-1]

            ref_dt = actual_entry or actual_exit
            if ref_dt is None:
                continue

            expected_entry = datetime.combine(ref_dt.date(), self._parse_time(first_period.start_time))
            expected_exit = datetime.combine(ref_dt.date(), self._parse_time(last_period.end_time))
            if expected_exit < expected_entry:
                expected_exit += timedelta(days=1)

            # Ajusta el ancla de fecha para minimizar diferencia cuando hay turnos nocturnos.
            if actual_entry and abs((actual_entry - expected_entry).total_seconds()) > 12 * 3600:
                if actual_entry > expected_entry:
                    expected_entry += timedelta(days=1)
                    expected_exit += timedelta(days=1)
                else:
                    expected_entry -= timedelta(days=1)
                    expected_exit -= timedelta(days=1)

            score = 0
            if actual_entry:
                score += abs((actual_entry - expected_entry).total_seconds())
            if actual_exit:
                adjusted_actual_exit = actual_exit
                if actual_entry and adjusted_actual_exit < actual_entry:
                    adjusted_actual_exit += timedelta(days=1)
                if not actual_entry and adjusted_actual_exit < expected_entry:
                    adjusted_actual_exit += timedelta(days=1)
                score += abs((adjusted_actual_exit - expected_exit).total_seconds())

            if best_score is None or score < best_score:
                best_score = score
                best_turn = turn

        return best_turn

    def _resolve_expected_interval(self, turn, date_key, actual_entry=None, actual_exit=None):
        if not turn or not getattr(turn, "periods", None):
            if actual_entry and actual_exit:
                start_dt = actual_entry.replace(minute=0, second=0, microsecond=0)
                end_dt = actual_exit.replace(minute=0, second=0, microsecond=0)
                if end_dt < start_dt:
                    end_dt += timedelta(days=1)
                return start_dt, end_dt
            if actual_entry:
                start_dt = actual_entry.replace(minute=0, second=0, microsecond=0)
                return start_dt, start_dt + timedelta(hours=8)
            if actual_exit:
                end_dt = actual_exit.replace(minute=0, second=0, microsecond=0)
                return end_dt - timedelta(hours=8), end_dt
            return None, None

        first_period = turn.periods[0]
        last_period = turn.periods[-1]
        start_dt = datetime.combine(date_key, self._parse_time(first_period.start_time))
        end_dt = datetime.combine(date_key, self._parse_time(last_period.end_time))
        if end_dt < start_dt:
            end_dt += timedelta(days=1)
        return start_dt, end_dt

    def _resolve_expected_interval_for_segment(
        self,
        turn,
        date_key,
        actual_entry=None,
        actual_exit=None,
        segment_index=0,
    ):
        if turn and getattr(turn, "periods", None) and segment_index < len(turn.periods):
            work_period = turn.periods[segment_index]
            start_dt = datetime.combine(date_key, self._parse_time(work_period.start_time))
            end_dt = datetime.combine(date_key, self._parse_time(work_period.end_time))
            if end_dt < start_dt:
                end_dt += timedelta(days=1)
            return start_dt, end_dt
        return self._resolve_expected_interval(turn, date_key, actual_entry, actual_exit)

    def _build_employee_segments(self, records):
        if not records:
            return []

        segments = []
        pending_entry = None

        for record in sorted(records, key=lambda item: item["date_time"]):
            if record["registration"] == 0:
                if pending_entry is not None:
                    segments.append({"entry": pending_entry, "exit": None})
                pending_entry = record
                continue

            if pending_entry is None:
                segments.append({"entry": None, "exit": record})
                continue

            segments.append({"entry": pending_entry, "exit": record})
            pending_entry = None

        if pending_entry is not None:
            segments.append({"entry": pending_entry, "exit": None})

        return segments

    def _parse_time(self, value):
        if isinstance(value, str):
            try:
                return datetime.strptime(value, "%H:%M:%S").time()
            except ValueError:
                return datetime.strptime(value, "%H:%M").time()
        return value

    def _normalize_registration(self, value):
        if value is None:
            return None

        if isinstance(value, bool):
            return 1 if value else 0

        if isinstance(value, (int, float)):
            parsed = int(value)
            if parsed in self.VALID_REGISTRATIONS:
                return parsed
            # CrossChex usa 3 como entrada y 4 como salida en algunos archivos.
            if parsed >= 0:
                if parsed == 3:
                    return 0
                if parsed == 4:
                    return 1
                return 0 if parsed % 2 == 0 else 1
            return None

        normalized = str(value).strip().lower()
        if not normalized:
            return None

        direct_values = {
            "0": 0,
            "1": 1,
            "2": 0,
            "3": 0,
            "4": 1,
            "5": 1,
            "entrada": 0,
            "ingreso": 0,
            "in": 0,
            "check in": 0,
            "clock in": 0,
            "salida": 1,
            "egreso": 1,
            "out": 1,
            "check out": 1,
            "clock out": 1,
        }

        if normalized in direct_values:
            return direct_values[normalized]

        digits_only = "".join(ch for ch in normalized if ch.isdigit())
        if digits_only:
            parsed = int(digits_only)
            if parsed in self.VALID_REGISTRATIONS:
                return parsed
            return 0 if parsed % 2 == 0 else 1

        return None

    def _compute_minutes_late(self, actual_entry, expected_entry, tolerance_minutes):
        if not actual_entry or not expected_entry:
            return 0
        delta = actual_entry - expected_entry
        return max(0, int(delta.total_seconds() // 60))

    def _compute_minutes_early(self, actual_exit, expected_exit, tolerance_minutes):
        if not actual_exit or not expected_exit:
            return 0
        delta = expected_exit - actual_exit
        minutes = max(0, int(delta.total_seconds() // 60))
        if minutes <= tolerance_minutes:
            return 0
        return minutes

    def _is_in_period(self, date_value, period):
        if not period:
            return True
        try:
            period_start = datetime.strptime(period + "-01", "%Y-%m-%d").date()
            next_month = (period_start.replace(day=28) + timedelta(days=4)).replace(day=1)
            return period_start <= date_value < next_month
        except ValueError:
            return True

import calendar
import re
from io import BytesIO
from datetime import date

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from openpyxl import Workbook

from src.attendance.api.schemas import MonthlyTurnPlanRequest
from src.attendance.database.connection import get_connection


router = APIRouter(prefix="/monthly-turns", tags=["Monthly Turn Assignments"])
PERIOD_PATTERN = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")


def _validate_period(period):
    if not PERIOD_PATTERN.fullmatch(period):
        raise HTTPException(status_code=400, detail="El período debe tener formato YYYY-MM.")


def _ensure_table(connection):
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS monthly_turn_assignments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            period TEXT NOT NULL,
            assignment_date TEXT NOT NULL,
            employee_number INTEGER NOT NULL,
            assignment_code TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (employee_number) REFERENCES employees(employee_number)
                ON DELETE CASCADE,
            UNIQUE (period, assignment_date, employee_number)
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS special_codes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT NOT NULL UNIQUE,
            description TEXT NOT NULL,
            active INTEGER NOT NULL DEFAULT 1
        )
        """
    )


def _period_dates(period):
    year, month = (int(value) for value in period.split("-"))
    return [date(year, month, day).isoformat() for day in range(1, calendar.monthrange(year, month)[1] + 1)]


def _excel_day(value):
    try:
        day = int(value)
    except (TypeError, ValueError):
        return None
    return day if 1 <= day <= 31 else None


def _read_monthly_turn_workbook(contents, period, employees, allowed_codes):
    workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(
        (
            index
            for index, row in enumerate(rows)
            if row and str(row[0] or '').strip().lower() == 'legajo'
        ),
        None,
    )
    if header_index is None:
        raise HTTPException(status_code=400, detail="El Excel debe tener una columna 'Legajo'.")

    header = rows[header_index]
    day_header = rows[header_index - 1] if header_index > 0 else ()
    normalized_headers = {
        str(value or '').strip().lower(): index
        for index, value in enumerate(header)
    }
    required_headers = {'legajo', 'empleado'}
    missing_headers = required_headers - normalized_headers.keys()
    area_header = 'área' if 'área' in normalized_headers else 'area'
    if area_header not in normalized_headers:
        missing_headers.add('área')
    if missing_headers:
        raise HTTPException(
            status_code=400,
            detail=f"Faltan columnas requeridas: {', '.join(sorted(missing_headers))}.",
        )
    employee_column = normalized_headers['legajo']
    day_columns = {
        day: index
        for index, value in enumerate(header)
        if (day := _excel_day(value)) is not None
    }
    day_columns.update({
        day: index
        for index, value in enumerate(day_header)
        if (day := _excel_day(value)) is not None
    })
    period_days = len(_period_dates(period))
    missing_days = [day for day in range(1, period_days + 1) if day not in day_columns]
    if missing_days:
        raise HTTPException(
            status_code=400,
            detail=f"Faltan columnas para los días: {', '.join(map(str, missing_days))}.",
        )

    employee_numbers = set(employees)
    assignments = []
    seen_employees = set()
    for row_number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not row or employee_column >= len(row) or row[employee_column] in (None, ''):
            continue
        try:
            employee_number = int(float(row[employee_column]))
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail=f"Fila {row_number}: legajo inválido.")
        if employee_number not in employee_numbers:
            raise HTTPException(status_code=400, detail=f"Fila {row_number}: no existe el legajo {employee_number}.")
        if employee_number in seen_employees:
            raise HTTPException(status_code=400, detail=f"El legajo {employee_number} aparece más de una vez.")
        seen_employees.add(employee_number)

        for day in range(1, period_days + 1):
            value = row[day_columns[day]] if day_columns[day] < len(row) else None
            code = str(value or '').strip().upper()
            if not code:
                continue
            if code not in allowed_codes:
                raise HTTPException(
                    status_code=400,
                    detail=f"Fila {row_number}, día {day}: el código {code} no existe o no está activo.",
                )
            assignments.append({
                'employee_number': employee_number,
                'assignment_date': f'{period}-{day:02d}',
                'assignment_code': code,
            })
    return assignments, len(seen_employees)


@router.get("/{period}")
def get_monthly_turn_plan(period: str):
    _validate_period(period)
    connection = get_connection()
    try:
        _ensure_table(connection)
        employees = connection.execute(
            """
            SELECT employee_number, name, sector, active
            FROM employees
            WHERE active = 1
            ORDER BY sector, employee_number
            """
        ).fetchall()
        rows = connection.execute(
            """
            SELECT employee_number, assignment_date, assignment_code
            FROM monthly_turn_assignments
            WHERE period = ?
            """,
            (period,),
        ).fetchall()
        assignments = {
            (row["employee_number"], row["assignment_date"]): row["assignment_code"]
            for row in rows
        }
        dates = _period_dates(period)
        return {
            "period": period,
            "dates": dates,
            "employees": [
                {
                    "employee_number": employee["employee_number"],
                    "name": employee["name"],
                    "sector": employee["sector"] or "Sin sector",
                    "active": bool(employee["active"]),
                    "assignments": {
                        assignment_date: assignments.get((employee["employee_number"], assignment_date), "")
                        for assignment_date in dates
                    },
                }
                for employee in employees
            ],
        }
    finally:
        connection.close()


@router.put("/{period}")
def save_monthly_turn_plan(period: str, payload: MonthlyTurnPlanRequest):
    _validate_period(period)
    valid_dates = set(_period_dates(period))
    connection = get_connection()
    try:
        _ensure_table(connection)
        allowed_codes = {
            row["code"]
            for row in connection.execute(
                """
                SELECT code FROM turns WHERE active = 1
                UNION
                SELECT code FROM special_codes WHERE active = 1
                """
            ).fetchall()
        }
        for assignment in payload.assignments:
            if assignment.assignment_date not in valid_dates:
                raise HTTPException(status_code=400, detail=f"La fecha {assignment.assignment_date} no pertenece al período.")
            employee_exists = connection.execute(
                "SELECT 1 FROM employees WHERE employee_number = ? AND active = 1",
                (assignment.employee_number,),
            ).fetchone()
            if employee_exists is None:
                raise HTTPException(status_code=400, detail=f"No existe el empleado {assignment.employee_number}.")
            assignment_code = assignment.assignment_code.strip().upper()
            if not assignment_code:
                raise HTTPException(status_code=400, detail="El código de asignación no puede estar vacío.")
            if assignment_code not in allowed_codes:
                raise HTTPException(
                    status_code=400,
                    detail=f"El código {assignment_code} no existe o no está activo como turno o código especial.",
                )

        connection.execute("DELETE FROM monthly_turn_assignments WHERE period = ?", (period,))
        connection.executemany(
            """
            INSERT INTO monthly_turn_assignments (
                period, assignment_date, employee_number, assignment_code
            ) VALUES (?, ?, ?, ?)
            """,
            [
                (period, item.assignment_date, item.employee_number, item.assignment_code.strip().upper())
                for item in payload.assignments
                if item.assignment_code.strip()
            ],
        )
        connection.commit()
        return {"period": period, "saved": len(payload.assignments)}
    except HTTPException:
        connection.rollback()
        raise
    except Exception as error:
        connection.rollback()
        raise HTTPException(status_code=400, detail=str(error))
    finally:
        connection.close()


@router.post("/{period}/import")
async def import_monthly_turn_plan(
    period: str,
    file: UploadFile = File(...),
):
    _validate_period(period)
    if not file.filename or not file.filename.lower().endswith('.xlsx'):
        raise HTTPException(status_code=400, detail='Debe subir un archivo Excel .xlsx válido.')

    connection = get_connection()
    try:
        _ensure_table(connection)
        employees = {
            row['employee_number']
            for row in connection.execute(
                'SELECT employee_number FROM employees WHERE active = 1'
            ).fetchall()
        }
        allowed_codes = {
            row['code']
            for row in connection.execute(
                """
                SELECT code FROM turns WHERE active = 1
                UNION
                SELECT code FROM special_codes WHERE active = 1
                """
            ).fetchall()
        }
        assignments, employee_count = _read_monthly_turn_workbook(
            await file.read(), period, employees, allowed_codes
        )
        connection.execute('DELETE FROM monthly_turn_assignments WHERE period = ?', (period,))
        connection.executemany(
            """
            INSERT INTO monthly_turn_assignments (
                period, assignment_date, employee_number, assignment_code
            ) VALUES (?, ?, ?, ?)
            """,
            [
                (period, item['assignment_date'], item['employee_number'], item['assignment_code'])
                for item in assignments
            ],
        )
        connection.commit()
        return {'period': period, 'saved': len(assignments), 'employees': employee_count}
    except HTTPException:
        connection.rollback()
        raise
    except Exception as error:
        connection.rollback()
        raise HTTPException(status_code=400, detail=f'No se pudo leer el Excel: {error}')
    finally:
        connection.close()


@router.get("/template/{period}")
def download_monthly_turn_template(period: str):
    _validate_period(period)
    connection = get_connection()
    workbook = Workbook()
    try:
        employees = connection.execute(
            """
            SELECT employee_number, name, sector
            FROM employees
            WHERE active = 1
            ORDER BY sector, employee_number
            """
        ).fetchall()
        sheet = workbook.active
        sheet.title = 'Planilla mensual'
        dates = _period_dates(period)
        sheet.append([f'PLANILLA MENSUAL {period}'])
        sheet.append([])
        sheet.append(['Legajo', 'Empleado', 'Área', *range(1, len(dates) + 1)])
        for employee in employees:
            sheet.append([
                employee['employee_number'],
                employee['name'],
                employee['sector'] or 'Sin sector',
                *([''] * len(dates)),
            ])
        sheet.freeze_panes = 'D4'

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        filename = f'planilla-turnos-{period}.xlsx'
        return StreamingResponse(
            output,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename="{filename}"'},
        )
    finally:
        connection.close()